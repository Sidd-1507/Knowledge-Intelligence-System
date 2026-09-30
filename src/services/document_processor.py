import os
import io
from langchain.docstore.document import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from pypdf import PdfReader
import docx
from src.utils.logger import get_logger

logger = get_logger(__name__)

class DocumentProcessor:
    def __init__(self, settings):
        self.chunk_size = settings.chunk_size
        self.chunk_overlap = settings.chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap
        )

    def process_file(self, file_content: bytes, filename: str) -> list[Document]:
        ext = os.path.splitext(filename)[1].lower()
        if ext == '.pdf':
            return self._process_pdf(file_content, filename)
        elif ext in ['.docx', '.doc']:
            return self._process_docx(file_content, filename)
        elif ext == '.txt':
            return self._process_txt(file_content, filename)
        else:
            logger.error(f"Unsupported file format: {ext}")
            raise ValueError(f"Unsupported file format: {ext}")

    def _process_pdf(self, content: bytes, filename: str) -> list[Document]:
        try:
            reader = PdfReader(io.BytesIO(content))
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return self._chunk_text(text, {"source": filename})
        except Exception as e:
            logger.error(f"Failed to process PDF {filename}: {e}")
            raise

    def _process_docx(self, content: bytes, filename: str) -> list[Document]:
        try:
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join([para.text for para in doc.paragraphs])
            return self._chunk_text(text, {"source": filename})
        except Exception as e:
            logger.error(f"Failed to process DOCX {filename}: {e}")
            raise

    def _process_txt(self, content: bytes, filename: str) -> list[Document]:
        try:
            text = content.decode('utf-8')
            return self._chunk_text(text, {"source": filename})
        except Exception as e:
            logger.error(f"Failed to process TXT {filename}: {e}")
            raise

    def _chunk_text(self, text: str, metadata: dict) -> list[Document]:
        chunks = self.text_splitter.split_text(text)
        documents = []
        total_chunks = len(chunks)
        for i, chunk in enumerate(chunks):
            chunk_metadata = metadata.copy()
            chunk_metadata.update({
                "chunk_index": i,
                "total_chunks": total_chunks
            })
            documents.append(Document(page_content=chunk, metadata=chunk_metadata))
        return documents
